from io import BytesIO
from tempfile import NamedTemporaryFile

import fitz
from docx import Document
from django.test import SimpleTestCase

from .certificate_extracter import extract_certificates
from .hackathon_extracter import extract_hackathons
from .internship_extracter import extract_internships
from .language_extracter import extract_languages, extract_tools
from .project_extracter import extract_projects
from .resume_parser import extract_resume_text
from .skill_extracter import extract_skills
from .skill_matcher import match_skills


class ResumeIntelligenceTests(SimpleTestCase):
    SAMPLE = """\
    HEMALATHA\n
    EXPERIENCE\n
    ServiceNow Virtual Intern | SmartBridge\n
    Duration: May 2026 - June 2026\n
    Machine Learning Intern | Coding Blocks\n
    Duration: June 2026 - July 2026\n

    PROJECTS\n
    CampusHire ATS | Python, Django, SQL\n
    Built resume parsing and ATS matching.\n
    Reclaim | React, TypeScript, Gemini\n
    Built a campus lost and found platform.\n

    EDUCATION\n
    B.Tech CSE\n

    TECHNICAL SKILLS\n
    Languages: Python, Java, C++\n
    Tools: Git, GitHub, Power BI, Excel, ServiceNow\n

    CERTIFICATIONS\n
    Deloitte Data Analytics | Forage\n

    HACKATHON\n
    BuildWise | Python, Flask, Gemini API\n
    """

    def test_docx_table_text_is_extracted(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "PROJECTS"
        table.cell(1, 0).text = "CampusHire | Python, Django, SQL"
        data = BytesIO()
        doc.save(data)
        with NamedTemporaryFile(suffix=".docx") as f:
            f.write(data.getvalue())
            f.flush()
            text = extract_resume_text(f.name)
        self.assertIn("CampusHire", text)
        self.assertIn("Python, Django, SQL", text)

    def test_pdf_text_is_extracted(self):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "PROJECTS\nCampusHire | Python, Django, SQL")
        data = doc.tobytes()
        doc.close()
        with NamedTemporaryFile(suffix=".pdf") as f:
            f.write(data)
            f.flush()
            text = extract_resume_text(f.name)
        self.assertIn("CampusHire", text)
        self.assertIn("Python, Django, SQL", text)

    def test_resume_sections_are_extracted(self):
        self.assertIn("CampusHire", extract_projects(self.SAMPLE))
        self.assertIn("ServiceNow Virtual Intern", extract_internships(self.SAMPLE))
        self.assertIn("Deloitte Data Analytics", extract_certificates(self.SAMPLE))
        self.assertIn("BuildWise", extract_hackathons(self.SAMPLE))
        self.assertIn("Python", extract_languages(self.SAMPLE))
        self.assertIn("Power BI", extract_tools(self.SAMPLE))

    def test_skills_and_ats_match_resume_content(self):
        skills = extract_skills(self.SAMPLE)
        self.assertIn("Python", skills)
        self.assertIn("Java", skills)
        self.assertIn("C++", skills)
        self.assertNotIn("C", skills)  # C++ must not be reduced to C.
        matched, missing, score = match_skills([], "Python, Django, SQL, React", candidate_text=self.SAMPLE)
        self.assertEqual(matched, ["Python", "Django", "SQL", "React"])
        self.assertEqual(missing, [])
        self.assertEqual(score, 100.0)
